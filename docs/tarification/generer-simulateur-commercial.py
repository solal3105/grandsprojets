#!/usr/bin/env python3
"""Génère le simulateur Excel de l'équipe commerciale : la grille de prix telle
qu'elle est en production sur /home2/tarification, un tableau de bord avec
graphiques, des scénarios à volonté, et des réglages que l'on peut modifier
pour voir ce que cela donne.

Ce classeur est distinct de `grille-de-prix-commerciaux.xlsx` (qui sert à
proposer une nouvelle grille et à nous la renvoyer) : ici on joue, on prépare
un rendez-vous, on compare des offres. Rien n'est à renvoyer.

Les formules reproduisent `home-src/src/v2/data/tarification.mjs` à
l'identique, fourchette publique comprise. Quand la grille du site change, on
met à jour le bloc de constantes ci-dessous et on relance :

    python3 docs/tarification/generer-simulateur-commercial.py

Le classeur sort à côté du script : `simulateur-tarification.xlsx`.
Dépendance : openpyxl (pip install openpyxl)."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------------------
# La grille en production : la copie exacte de tarification.mjs (10/09/2026).
# ---------------------------------------------------------------------------

ANCRE = {"population": 12000, "prix": 200}
EXPOSANT = 0.5
POPULATION = {"min": 500, "max": 2500000}
# Dans l'ordre de la page du site, avec la couleur de chaque module
MODULES = [
    ("carte", "Carte des projets urbains", 1, "C4002A"),
    ("travaux", "Travaux du quotidien", 0.6, "B45309"),
    ("chantiers", "Chantiers et arrêtés", 3, "0B7A4A"),
    ("participer", "Signalement", 2, "1B5FA8"),
    ("diagnostic", "Diagnostic terrain", 0.8, "7546CC"),
]
CHANTIERS_DEMI_POIDS = {"sous": 5000, "plein": 20000}
MISE_EN_SERVICE = {"mois": 6, "offerteSous": 0}
REMISE_PAR_MODULE_AJOUTE = 0.10
ENGAGEMENTS = [(1, 0), (2, 0.10), (3, 0.15), (4, 0.20)]
FOURCHETTE_MARGE = 0.20
SEUILS = [
    (60000, "Procédure adaptée"),
    (90000, "Publicité obligatoire"),
    (216000, "Procédure formalisée"),
]
COURBE = [500, 800, 1500, 3000, 5000, 8000, 12000, 20000, 35000, 50000, 80000, 150000, 300000, 500000, 1000000]

# ---------------------------------------------------------------------------
# Mise en forme
# ---------------------------------------------------------------------------

ROUGE = "FF0037"
ENCRE = "111111"
GRIS_TEXTE = "6B7280"
FILL_SAISIE = PatternFill("solid", fgColor="FFF4D6")
FILL_PROD = PatternFill("solid", fgColor="F1F2F4")
FILL_ENTETE = PatternFill("solid", fgColor=ENCRE)
FILL_TUILE = PatternFill("solid", fgColor="FAFAFA")
FILL_TUILE_FORTE = PatternFill("solid", fgColor=ENCRE)
FONT_TITRE = Font(name="Calibri", size=20, bold=True, color=ENCRE)
FONT_SOUS_TITRE = Font(name="Calibri", size=11, color=GRIS_TEXTE)
FONT_ENTETE = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_SECTION = Font(name="Calibri", size=12, bold=True, color=ENCRE)
FONT_GRAS = Font(name="Calibri", size=11, bold=True, color=ENCRE)
FONT_NORMAL = Font(name="Calibri", size=11, color=ENCRE)
FONT_NOTE = Font(name="Calibri", size=9, color=GRIS_TEXTE)
FONT_TUILE_LIBELLE = Font(name="Calibri", size=9, bold=True, color=GRIS_TEXTE)
FONT_TUILE_VALEUR = Font(name="Calibri", size=22, bold=True, color=ENCRE)
FONT_TUILE_LIBELLE_CLAIR = Font(name="Calibri", size=9, bold=True, color="D1D5DB")
FONT_TUILE_VALEUR_CLAIR = Font(name="Calibri", size=22, bold=True, color="FFFFFF")
FONT_TUILE_NOTE_CLAIR = Font(name="Calibri", size=9, color="D1D5DB")
BORDURE = Border(*(Side(style="thin", color="E5E7EB"),) * 4)
BORDURE_TUILE = Border(*(Side(style="medium", color="E5E7EB"),) * 4)
GAUCHE = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)
DROITE = Alignment(horizontal="right", vertical="center")
HAUT = Alignment(horizontal="left", vertical="top", wrap_text=True)

FMT_EUROS = '#,##0" €"'
FMT_EUROS_SIGNE = '+#,##0" €";-#,##0" €";0" €"'
FMT_POURCENT = "0%"
FMT_ENTIER = "#,##0"
FMT_DECIMAL = "0.00"
DEVERROUILLE = Protection(locked=False)

wb = Workbook()


def nom_defini(nom, feuille, plage):
    wb.defined_names[nom] = DefinedName(nom, attr_text=f"'{feuille}'!{plage}")


def entete(ws, ligne, colonnes, debut=1, hauteur=None):
    for i, texte in enumerate(colonnes):
        c = ws.cell(row=ligne, column=debut + i, value=texte)
        c.fill, c.font, c.alignment, c.border = FILL_ENTETE, FONT_ENTETE, CENTRE, BORDURE
    if hauteur:
        ws.row_dimensions[ligne].height = hauteur


def cellule(ws, ligne, colonne, valeur=None, fmt=None, fill=None, font=FONT_NORMAL, align=None, saisie=False, bordure=True):
    c = ws.cell(row=ligne, column=colonne, value=valeur)
    c.font = font
    if bordure:
        c.border = BORDURE
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if saisie:
        c.fill = FILL_SAISIE
        c.protection = DEVERROUILLE
    return c


def largeurs(ws, valeurs, debut=1):
    for i, largeur in enumerate(valeurs):
        ws.column_dimensions[get_column_letter(debut + i)].width = largeur


def proteger(ws):
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def titre(ws, texte, sous_titre=None, colonne=2):
    ws.cell(row=2, column=colonne, value=texte).font = FONT_TITRE
    if sous_titre:
        c = ws.cell(row=3, column=colonne, value=sous_titre)
        c.font, c.alignment = FONT_SOUS_TITRE, GAUCHE


# ---------------------------------------------------------------------------
# Le calcul, écrit une fois. `a` : adresse de chaque intermédiaire ; `sel` :
# adresse de la case Oui / Non de chaque module (absente = toujours retenu) ;
# `n` : expression qui compte les modules retenus ; `annees` et `population` :
# adresses ou valeurs.
# ---------------------------------------------------------------------------

ORDRE = ["pb", "unite", "pch"] + [f"prix_{c}" for c, _, _, _ in MODULES] + [
    "brut", "n", "taux_mod", "rem_mod", "apres", "taux_eng", "rem_eng", "mensuel", "annuel", "setup", "total", "seuil",
    "pas_m", "bas_m", "haut_m", "pas_t", "bas_t", "haut_t",
]
LIBELLES = {
    "pb": "Population retenue", "unite": "Prix d'une unité de poids", "pch": "Poids chantiers",
    "brut": "Total avant remises", "n": "Modules retenus", "taux_mod": "Taux remise multi-modules",
    "rem_mod": "Remise multi-modules", "apres": "Après remise multi-modules", "taux_eng": "Taux engagement",
    "rem_eng": "Remise engagement", "mensuel": "Mensuel", "annuel": "Annuel", "setup": "Mise en service",
    "total": "Total sur la durée", "seuil": "Seuil", "pas_m": "Pas d'arrondi (mensuel)", "bas_m": "Fourchette basse (mensuel)",
    "haut_m": "Fourchette haute (mensuel)", "pas_t": "Pas d'arrondi (total)", "bas_t": "Fourchette basse (total)",
    "haut_t": "Fourchette haute (total)",
}
FORMATS = {
    "pb": FMT_ENTIER, "unite": FMT_EUROS, "pch": FMT_DECIMAL, "brut": FMT_EUROS, "n": "0", "taux_mod": FMT_POURCENT,
    "rem_mod": FMT_EUROS, "apres": FMT_EUROS, "taux_eng": FMT_POURCENT, "rem_eng": FMT_EUROS, "mensuel": FMT_EUROS,
    "annuel": FMT_EUROS, "setup": FMT_EUROS, "total": FMT_EUROS, "seuil": "@", "pas_m": FMT_ENTIER, "bas_m": FMT_EUROS,
    "haut_m": FMT_EUROS, "pas_t": FMT_ENTIER, "bas_t": FMT_EUROS, "haut_t": FMT_EUROS,
}
for cle, nom, _, _ in MODULES:
    LIBELLES[f"prix_{cle}"] = f"Prix {nom}"
    FORMATS[f"prix_{cle}"] = FMT_EUROS


def pas_arrondi(v):
    return f"IF({v}<1000,10,IF({v}<10000,100,IF({v}<100000,1000,10000)))"


def formules(a, sel, n, annees, population):
    f = {}
    f["pb"] = f"=MIN(Population_maximum,MAX(Population_minimum,ROUND({population},0)))"
    f["unite"] = f"=Prix_ancre*({a['pb']}/Population_ancre)^Exposant"
    f["pch"] = (
        f"=IF({a['pb']}<=Chantiers_demi_poids_sous,Poids_chantiers/2,"
        f"IF({a['pb']}>=Chantiers_plein_poids_des,Poids_chantiers,"
        f"Poids_chantiers*(0.5+0.5*(LN({a['pb']})-LN(Chantiers_demi_poids_sous))"
        f"/(LN(Chantiers_plein_poids_des)-LN(Chantiers_demi_poids_sous)))))"
    )
    for cle, _, _, _ in MODULES:
        poids = a["pch"] if cle == "chantiers" else f"Poids_{cle}"
        prix = f"{a['unite']}*{poids}"
        f[f"prix_{cle}"] = f"=IF({sel[cle]}=\"Oui\",{prix},0)" if sel.get(cle) else f"={prix}"
    f["brut"] = f"=SUM({a['prix_range']})"
    f["n"] = f"={n}"
    f["taux_mod"] = f"=MIN(0.9,MAX(0,{a['n']}-1)*Remise_par_module_ajoute)"
    f["rem_mod"] = f"=MAX({a['prix_range']})*{a['taux_mod']}"
    f["apres"] = f"={a['brut']}-{a['rem_mod']}"
    f["taux_eng"] = f"=IFERROR(VLOOKUP({annees},Table_engagements,4,FALSE),0)"
    f["rem_eng"] = f"={a['apres']}*{a['taux_eng']}"
    f["mensuel"] = f"={a['apres']}-{a['rem_eng']}"
    f["annuel"] = f"={a['mensuel']}*12"
    f["setup"] = f"=IF({a['pb']}<Mise_en_service_offerte_sous,0,{a['mensuel']}*Mise_en_service_mois)"
    f["total"] = f"={a['setup']}+{a['annuel']}*IF(N({annees})>0,{annees},1)"
    f["seuil"] = (
        f"=IF({a['total']}>=Seuil_3,\"{SEUILS[2][1]}\","
        f"IF({a['total']}>=Seuil_2,\"{SEUILS[1][1]}\","
        f"IF({a['total']}>=Seuil_1,\"{SEUILS[0][1]}\",\"Sous les seuils\")))"
    )
    f["pas_m"] = f"={pas_arrondi(a['mensuel'])}"
    f["bas_m"] = f"=FLOOR({a['mensuel']}*(1-Fourchette_marge),{a['pas_m']})"
    f["haut_m"] = f"=CEILING({a['mensuel']}*(1+Fourchette_marge),{a['pas_m']})"
    f["pas_t"] = f"={pas_arrondi(a['total'])}"
    f["bas_t"] = f"=FLOOR({a['total']}*(1-Fourchette_marge),{a['pas_t']})"
    f["haut_t"] = f"=CEILING({a['total']}*(1+Fourchette_marge),{a['pas_t']})"
    return f


def bloc_horizontal(ws, ligne, colonne_debut, sel, n, annees, population):
    """Pose le calcul sur une ligne, une colonne par intermédiaire. Rend les adresses."""
    a = {}
    for i, cle in enumerate(ORDRE):
        a[cle] = f"{get_column_letter(colonne_debut + i)}{ligne}"
    premier = ORDRE.index("prix_carte")
    a["prix_range"] = (
        f"{get_column_letter(colonne_debut + premier)}{ligne}:"
        f"{get_column_letter(colonne_debut + premier + len(MODULES) - 1)}{ligne}"
    )
    f = formules(a, sel, n, annees, population)
    for i, cle in enumerate(ORDRE):
        c = ws.cell(row=ligne, column=colonne_debut + i, value=f[cle])
        c.number_format, c.font = FORMATS[cle], FONT_NOTE
    return a


def entetes_bloc(ws, ligne, colonne_debut):
    for i, cle in enumerate(ORDRE):
        c = ws.cell(row=ligne, column=colonne_debut + i, value=LIBELLES[cle])
        c.font, c.alignment = FONT_NOTE, CENTRE
        ws.column_dimensions[get_column_letter(colonne_debut + i)].width = 13


# ---------------------------------------------------------------------------
# Feuille Réglages : la grille de production, et la valeur utilisée ici
# ---------------------------------------------------------------------------

REGLAGES = "Réglages"
ws_r = wb.create_sheet(REGLAGES)
ws_r.sheet_view.showGridLines = False
largeurs(ws_r, [3, 48, 22, 22, 12, 56])
titre(ws_r, "Les réglages de la grille", (
    "La colonne grise est la grille en production sur le site. La colonne jaune est celle que tout le classeur "
    "utilise : elle part de la production, et vous pouvez la modifier pour voir ce que cela change. Pour revenir "
    "à la grille du site, recopiez la valeur grise."
))
ws_r.merge_cells(start_row=3, start_column=2, end_row=3, end_column=6)
ws_r.row_dimensions[3].height = 44

ligne_r = 5


def section_r(texte, explication=None):
    global ligne_r
    ligne_r += 1
    ws_r.cell(row=ligne_r, column=2, value=texte).font = FONT_SECTION
    if explication:
        ligne_r += 1
        c = ws_r.cell(row=ligne_r, column=2, value=explication)
        c.font, c.alignment = FONT_NOTE, GAUCHE
        ws_r.merge_cells(start_row=ligne_r, start_column=2, end_row=ligne_r, end_column=6)
        ws_r.row_dimensions[ligne_r].height = 15 * (len(explication) // 110 + 1) + 4
    ligne_r += 1
    entete(ws_r, ligne_r, ["Réglage", "Sur le site", "Utilisé ici", "Écart", "Ce que cela veut dire"], debut=2)


def reglage(libelle, nom, valeur, fmt, explication=""):
    global ligne_r
    ligne_r += 1
    cellule(ws_r, ligne_r, 2, libelle, align=GAUCHE)
    cellule(ws_r, ligne_r, 3, valeur, fmt=fmt, fill=FILL_PROD, align=DROITE)
    cellule(ws_r, ligne_r, 4, valeur, fmt=fmt, align=DROITE, saisie=True)
    cellule(ws_r, ligne_r, 5, f'=IF(D{ligne_r}=C{ligne_r},"","modifié")', align=CENTRE, font=FONT_NOTE)
    cellule(ws_r, ligne_r, 6, explication, align=GAUCHE, font=FONT_NOTE)
    nom_defini(nom, REGLAGES, f"$D${ligne_r}")
    return ligne_r


section_r("La courbe", (
    "Le prix d'une unité de poids suit la population en puissance. Deux chiffres suffisent : le prix à l'ancre et "
    "l'exposant. Avec un exposant de 0,5, une commune dix fois plus peuplée paie un peu plus de trois fois plus."
))
reglage("Population de l'ancre", "Population_ancre", ANCRE["population"], FMT_ENTIER, "La commune de référence de la courbe")
reglage("Prix mensuel HT d'une unité de poids à l'ancre", "Prix_ancre", ANCRE["prix"], FMT_EUROS, "Le prix du module Carte pour cette commune")
reglage("Exposant de la courbe", "Exposant", EXPOSANT, FMT_DECIMAL, "Plus il est haut, plus les grandes villes paient")
reglage("Population minimale", "Population_minimum", POPULATION["min"], FMT_ENTIER, "En dessous, la page ramène à ce plancher")
reglage("Population maximale", "Population_maximum", POPULATION["max"], FMT_ENTIER, "Au-dessus, la page ramène à ce plafond")

section_r("Le poids de chaque module", "Le prix d'un module est le prix de l'unité multiplié par son poids.")
for cle, nom, poids, _ in MODULES:
    reglage(nom, f"Poids_{cle}", poids, FMT_DECIMAL, "")

section_r("Le module Chantiers et arrêtés dans les petites communes", (
    "Il vaut la moitié de son poids jusqu'au premier seuil, son poids entier à partir du second, et monte en pente "
    "douce entre les deux."
))
reglage("Demi-poids jusqu'à", "Chantiers_demi_poids_sous", CHANTIERS_DEMI_POIDS["sous"], FMT_ENTIER, "habitants")
reglage("Poids entier à partir de", "Chantiers_plein_poids_des", CHANTIERS_DEMI_POIDS["plein"], FMT_ENTIER, "habitants")

section_r("Les remises", (
    "Prendre plusieurs modules fait baisser le plus cher d'entre eux, de ce taux par module ajouté au premier. "
    "L'engagement baisse ensuite tout l'abonnement, chaque année."
))
reglage("Remise par module ajouté, sur le module le plus cher", "Remise_par_module_ajoute", REMISE_PAR_MODULE_AJOUTE, FMT_POURCENT, "Deux modules : ce taux, trois : le double, et ainsi de suite")
ligne_eng_debut = ligne_r + 1
for annees, remise in ENGAGEMENTS:
    ligne = reglage(f"Engagement {annees} an{'s' if annees > 1 else ''}", f"Remise_engagement_{annees}", remise, FMT_POURCENT, "Remise sur l'abonnement, chaque année")
    ws_r.cell(row=ligne, column=1, value=annees).font = FONT_NOTE
nom_defini("Table_engagements", REGLAGES, f"$A${ligne_eng_debut}:$D${ligne_r}")

section_r("La mise en service", (
    "Facturée une seule fois, la première année, en mois d'abonnement au tarif obtenu. Un seuil d'habitants à zéro "
    "veut dire qu'elle n'est jamais offerte."
))
reglage("Mise en service, en mois d'abonnement", "Mise_en_service_mois", MISE_EN_SERVICE["mois"], "0", "")
reglage("Mise en service offerte sous", "Mise_en_service_offerte_sous", MISE_EN_SERVICE["offerteSous"], FMT_ENTIER, "habitants, zéro = jamais offerte")

section_r("La fourchette affichée au public", (
    "Un visiteur du site ne voit pas le prix exact : une fourchette autour du prix, à plus ou moins cette marge, "
    "arrondie vers l'extérieur (à la dizaine sous 1 000 euros, à la centaine sous 10 000, au millier sous 100 000, "
    "à la dizaine de milliers au-delà). Vous voyez les deux dans ce classeur."
))
reglage("Marge de la fourchette", "Fourchette_marge", FOURCHETTE_MARGE, FMT_POURCENT, "De part et d'autre du prix exact")

section_r("Les seuils de la commande publique", (
    "Ils s'apprécient sur le total hors taxes de toute la durée, mise en service comprise. Ils ne se négocient pas : "
    "ils sont là pour lire le résultat."
))
for i, (montant, nom) in enumerate(SEUILS, start=1):
    reglage(nom, f"Seuil_{i}", montant, FMT_EUROS, "hors taxes, sur toute la durée")

ligne_r += 2
c = ws_r.cell(row=ligne_r, column=2, value=(
    "Un mot « modifié » dans la colonne Écart signale que la valeur utilisée n'est plus celle du site. Le tableau de "
    "bord affiche le nombre de réglages modifiés."
))
c.font, c.alignment = FONT_NOTE, GAUCHE
ws_r.merge_cells(start_row=ligne_r, start_column=2, end_row=ligne_r, end_column=6)
PLAGE_ECARTS = f"'{REGLAGES}'!$E$1:$E${ligne_r}"
ws_r.freeze_panes = "A5"
proteger(ws_r)

# ---------------------------------------------------------------------------
# Feuille Grille : la courbe sur toutes les tailles de communes
# ---------------------------------------------------------------------------

GRILLE = "Grille"
ws_g = wb.create_sheet(GRILLE)
ws_g.sheet_view.showGridLines = False
titre(ws_g, "La grille sur toutes les tailles de communes", (
    "Prix mensuels hors taxes, avec la durée d'engagement choisie ci-dessous : chaque module pris seul, puis la suite "
    "complète des cinq modules avec sa remise multi-modules. Les populations se modifient."
))
ws_g.merge_cells(start_row=3, start_column=2, end_row=3, end_column=12)
ws_g.row_dimensions[3].height = 32
cellule(ws_g, 5, 2, "Durée d'engagement appliquée à cette grille, en années", align=GAUCHE)
valid_annees_g = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=False)
ws_g.add_data_validation(valid_annees_g)
c = cellule(ws_g, 5, 3, 1, fmt="0", align=CENTRE, saisie=True)
valid_annees_g.add(c)
ANNEES_GRILLE = f"'{GRILLE}'!$C$5"

colonnes_g = ["Habitants", "Une unité de poids"] + [f"{nom} seul" for _, nom, _, _ in MODULES] + [
    "Les 5 modules, par mois", "Mise en service", "Total sur la durée", "Seuil dépassé", "Fourchette publique, par mois",
]
entete(ws_g, 7, colonnes_g, debut=2, hauteur=46)
largeurs(ws_g, [3, 12, 14] + [16] * len(MODULES) + [16, 16, 18, 20, 26])
COL_BLOC_G = 2 + len(colonnes_g) + 2
entetes_bloc(ws_g, 7, COL_BLOC_G)
ws_g.column_dimensions[get_column_letter(COL_BLOC_G - 1)].width = 3
for i in range(len(ORDRE)):
    col = get_column_letter(COL_BLOC_G + i)
    ws_g.column_dimensions[col].hidden = True
    ws_g.column_dimensions[col].outlineLevel = 1

LIGNE_G0 = 8
for i, pop in enumerate(COURBE):
    r = LIGNE_G0 + i
    cellule(ws_g, r, 2, pop, fmt=FMT_ENTIER, align=DROITE, saisie=True)
    a = bloc_horizontal(ws_g, r, COL_BLOC_G, {}, len(MODULES), ANNEES_GRILLE, f"$B{r}")
    cellule(ws_g, r, 3, f"={a['unite']}", fmt=FMT_EUROS, align=DROITE)
    for j, (cle, _, _, _) in enumerate(MODULES):
        cellule(ws_g, r, 4 + j, f"={a[f'prix_{cle}']}*(1-{a['taux_eng']})", fmt=FMT_EUROS, align=DROITE)
    col = 4 + len(MODULES)
    cellule(ws_g, r, col, f"={a['mensuel']}", fmt=FMT_EUROS, align=DROITE, font=FONT_GRAS)
    cellule(ws_g, r, col + 1, f"={a['setup']}", fmt=FMT_EUROS, align=DROITE)
    cellule(ws_g, r, col + 2, f"={a['total']}", fmt=FMT_EUROS, align=DROITE)
    cellule(ws_g, r, col + 3, f"={a['seuil']}", align=CENTRE, font=FONT_NOTE)
    cellule(ws_g, r, col + 4, f'=TRIM(TEXT({a["bas_m"]},"# ##0"))&" à "&TRIM(TEXT({a["haut_m"]},"# ##0"))&" €"', align=CENTRE, font=FONT_NOTE)
LIGNE_G1 = LIGNE_G0 + len(COURBE) - 1
COL_SUITE_G = 4 + len(MODULES)
ws_g.freeze_panes = "D8"
ws_g.sheet_properties.outlinePr.summaryRight = False
proteger(ws_g)

# ---------------------------------------------------------------------------
# Feuille Scénarios : autant de cas que l'on veut
# ---------------------------------------------------------------------------

SCENARIOS = "Scénarios"
ws_s = wb.create_sheet(SCENARIOS)
ws_s.sheet_view.showGridLines = False
titre(ws_s, "Vos scénarios", (
    "Une ligne par cas : un nom, la population, Oui ou Non pour chaque module, la durée d'engagement. Le prix se "
    "calcule tout seul, avec la fourchette que le visiteur verrait sur le site. Les exemples se remplacent."
))
ws_s.merge_cells(start_row=3, start_column=2, end_row=3, end_column=16)
ws_s.row_dimensions[3].height = 32

colonnes_s = ["Scénario", "Habitants"] + [nom for _, nom, _, _ in MODULES] + ["Engagement (années)"] + [
    "Mensuel HT", "Annuel HT", "Mise en service", "Total sur la durée", "Seuil dépassé", "Fourchette publique, par mois",
    "Fourchette publique, total",
]
entete(ws_s, 5, colonnes_s, debut=2, hauteur=46)
largeurs(ws_s, [3, 30, 11] + [13] * len(MODULES) + [12, 13, 13, 14, 16, 20, 24, 26])
ws_s.freeze_panes = "D6"

valid_oui_non_s = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
valid_annees_s = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=True)
ws_s.add_data_validation(valid_oui_non_s)
ws_s.add_data_validation(valid_annees_s)

EXEMPLES = [
    ("Village, la carte seule", 800, ["carte"], 3),
    ("Bourg, carte et travaux", 3000, ["carte", "travaux"], 3),
    ("Bourg, chantiers seul", 3000, ["chantiers"], 1),
    ("Petite ville, carte, travaux et signalement", 12000, ["carte", "travaux", "participer"], 3),
    ("Petite ville, la suite complète", 12000, ["carte", "travaux", "chantiers", "participer", "diagnostic"], 4),
    ("Ville moyenne, carte et signalement", 50000, ["carte", "participer"], 2),
    ("Ville moyenne, la suite complète", 50000, ["carte", "travaux", "chantiers", "participer", "diagnostic"], 3),
    ("Grande ville, carte et diagnostic", 150000, ["carte", "diagnostic"], 3),
    ("Métropole, la carte seule", 500000, ["carte"], 4),
]
NB_SC = 25
COL_ENG_S = 4 + len(MODULES)
COL_SORTIE_S = COL_ENG_S + 1
COL_BLOC_S = COL_SORTIE_S + 7 + 1
entetes_bloc(ws_s, 5, COL_BLOC_S)
ws_s.column_dimensions[get_column_letter(COL_BLOC_S - 1)].width = 3
for i in range(len(ORDRE)):
    col = get_column_letter(COL_BLOC_S + i)
    ws_s.column_dimensions[col].hidden = True
    ws_s.column_dimensions[col].outlineLevel = 1

LIGNE_S0 = 6
for i in range(NB_SC):
    r = LIGNE_S0 + i
    ex = EXEMPLES[i] if i < len(EXEMPLES) else None
    cellule(ws_s, r, 2, ex[0] if ex else None, align=GAUCHE, saisie=True)
    cellule(ws_s, r, 3, ex[1] if ex else None, fmt=FMT_ENTIER, align=DROITE, saisie=True)
    sel = {}
    for j, (cle, _, _, _) in enumerate(MODULES):
        c = cellule(ws_s, r, 4 + j, (None if not ex else ("Oui" if cle in ex[2] else "Non")), align=CENTRE, saisie=True)
        valid_oui_non_s.add(c)
        sel[cle] = f"${get_column_letter(4 + j)}{r}"
    c = cellule(ws_s, r, COL_ENG_S, ex[3] if ex else None, fmt="0", align=CENTRE, saisie=True)
    valid_annees_s.add(c)
    annees_ref = f"${get_column_letter(COL_ENG_S)}{r}"
    sel_range = f"$D{r}:${get_column_letter(3 + len(MODULES))}{r}"
    a = bloc_horizontal(ws_s, r, COL_BLOC_S, sel, f'COUNTIF({sel_range},"Oui")', annees_ref, f"$C{r}")
    vide = f'$C{r}=""'
    sorties = [
        (f'=IF({vide},"",{a["mensuel"]})', FMT_EUROS, FONT_GRAS, DROITE),
        (f'=IF({vide},"",{a["annuel"]})', FMT_EUROS, FONT_NORMAL, DROITE),
        (f'=IF({vide},"",{a["setup"]})', FMT_EUROS, FONT_NORMAL, DROITE),
        (f'=IF({vide},"",{a["total"]})', FMT_EUROS, FONT_GRAS, DROITE),
        (f'=IF({vide},"",{a["seuil"]})', "@", FONT_NOTE, CENTRE),
        (f'=IF({vide},"",TRIM(TEXT({a["bas_m"]},"# ##0"))&" à "&TRIM(TEXT({a["haut_m"]},"# ##0"))&" €")', "@", FONT_NOTE, CENTRE),
        (f'=IF({vide},"",TRIM(TEXT({a["bas_t"]},"# ##0"))&" à "&TRIM(TEXT({a["haut_t"]},"# ##0"))&" €")', "@", FONT_NOTE, CENTRE),
    ]
    for k, (formule, fmt, font, align) in enumerate(sorties):
        cellule(ws_s, r, COL_SORTIE_S + k, formule, fmt=fmt, font=font, align=align)
LIGNE_S1 = LIGNE_S0 + NB_SC - 1
# Le seuil dépassé se colore, comme la jauge de la page du site
col_seuil_s = get_column_letter(COL_SORTIE_S + 4)
plage_seuil_s = f"{col_seuil_s}{LIGNE_S0}:{col_seuil_s}{LIGNE_S1}"
ws_s.conditional_formatting.add(plage_seuil_s, CellIsRule(operator="equal", formula=['"Sous les seuils"'], font=Font(color="15803D", size=9)))
ws_s.conditional_formatting.add(plage_seuil_s, CellIsRule(operator="equal", formula=[f'"{SEUILS[2][1]}"'], font=Font(color="B91C1C", size=9, bold=True)))
ws_s.conditional_formatting.add(plage_seuil_s, CellIsRule(operator="equal", formula=[f'"{SEUILS[1][1]}"'], font=Font(color="B45309", size=9)))
ws_s.conditional_formatting.add(plage_seuil_s, CellIsRule(operator="equal", formula=[f'"{SEUILS[0][1]}"'], font=Font(color="B45309", size=9)))
ws_s.sheet_properties.outlinePr.summaryRight = False
proteger(ws_s)

# ---------------------------------------------------------------------------
# Feuille Calculs (masquée) : le scénario du tableau de bord, et le même
# scénario sur les quatre durées d'engagement pour le graphique
# ---------------------------------------------------------------------------

CALCULS = "Calculs"
ws_c = wb.create_sheet(CALCULS)
ws_c.sheet_state = "hidden"
ws_c.cell(row=1, column=1, value="Feuille de calcul du tableau de bord. Ne rien saisir ici.").font = FONT_NOTE
entetes_bloc(ws_c, 3, 3)
ws_c.cell(row=3, column=1, value="Cas").font = FONT_NOTE
ws_c.cell(row=3, column=2, value="Années").font = FONT_NOTE

TB = "Tableau de bord"
POP_TB = f"'{TB}'!$C$7"
ANNEES_TB = f"'{TB}'!$C$13"
SEL_TB = {cle: f"'{TB}'!$C${8 + i}" for i, (cle, _, _, _) in enumerate(MODULES)}
SEL_RANGE_TB = f"'{TB}'!$C$8:$C${7 + len(MODULES)}"
N_TB = f'COUNTIF({SEL_RANGE_TB},"Oui")'

ws_c.cell(row=4, column=1, value="Scénario du tableau de bord").font = FONT_NOTE
ws_c.cell(row=4, column=2, value=f"={ANNEES_TB}").font = FONT_NOTE
A_TB = bloc_horizontal(ws_c, 4, 3, SEL_TB, N_TB, ANNEES_TB, POP_TB)
# Les mêmes adresses, vues depuis une autre feuille
A_TBX = {k: f"{CALCULS}!{v}" for k, v in A_TB.items()}
LIGNES_DUREES = {}
for annees, _ in ENGAGEMENTS:
    r = 5 + annees
    ws_c.cell(row=r, column=1, value=f"Même scénario sur {annees} an{'s' if annees > 1 else ''}").font = FONT_NOTE
    ws_c.cell(row=r, column=2, value=annees).font = FONT_NOTE
    LIGNES_DUREES[annees] = bloc_horizontal(ws_c, r, 3, SEL_TB, N_TB, f"$B${r}", POP_TB)

# Les séries du graphique par durée : libellé, mensuel, total
ws_c.cell(row=12, column=1, value="Durée").font = FONT_NOTE
ws_c.cell(row=12, column=2, value="Abonnement mensuel").font = FONT_NOTE
ws_c.cell(row=12, column=3, value="Total sur la durée").font = FONT_NOTE
for annees, _ in ENGAGEMENTS:
    r = 12 + annees
    ws_c.cell(row=r, column=1, value=f"{annees} an{'s' if annees > 1 else ''}")
    ws_c.cell(row=r, column=2, value=f"={LIGNES_DUREES[annees]['mensuel']}").number_format = FMT_EUROS
    ws_c.cell(row=r, column=3, value=f"={LIGNES_DUREES[annees]['total']}").number_format = FMT_EUROS

# Les séries du graphique par module : nom, prix retenu dans le scénario
ws_c.cell(row=19, column=1, value="Module").font = FONT_NOTE
ws_c.cell(row=19, column=2, value="Prix mensuel dans ce scénario").font = FONT_NOTE
for i, (cle, nom, _, _) in enumerate(MODULES):
    r = 20 + i
    ws_c.cell(row=r, column=1, value=nom)
    # Le module le plus cher porte la remise multi-modules : son prix affiché en tient compte
    ws_c.cell(row=r, column=2, value=(
        f"=IF({A_TB[f'prix_{cle}']}=0,0,IF(AND({A_TB['taux_mod']}>0,{A_TB[f'prix_{cle}']}=MAX({A_TB['prix_range']})),"
        f"{A_TB[f'prix_{cle}']}*(1-{A_TB['taux_mod']}),{A_TB[f'prix_{cle}']}))"
    )).number_format = FMT_EUROS

# ---------------------------------------------------------------------------
# Feuille Tableau de bord
# ---------------------------------------------------------------------------

ws_t = wb.create_sheet(TB, 0)
ws_t.sheet_view.showGridLines = False
ws_t.sheet_view.zoomScale = 110
largeurs(ws_t, [2, 30, 14, 3, 17, 17, 17, 17, 3, 14, 14, 14, 14, 14, 14, 14, 14])
ws_t.cell(row=2, column=2, value="Simulateur de tarification Open Projets").font = FONT_TITRE
c = ws_t.cell(row=3, column=2, value=(
    "Réglez la commune dans les cases jaunes : tout le tableau de bord suit. Les montants sont hors taxes."
))
c.font, c.alignment = FONT_SOUS_TITRE, GAUCHE
ws_t.merge_cells(start_row=3, start_column=2, end_row=3, end_column=8)

# La commune et son choix
ws_t.cell(row=5, column=2, value="Votre scénario").font = FONT_SECTION
valid_oui_non_t = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=False)
valid_annees_t = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=False)
ws_t.add_data_validation(valid_oui_non_t)
ws_t.add_data_validation(valid_annees_t)
cellule(ws_t, 6, 2, "Collectivité", align=GAUCHE)
cellule(ws_t, 6, 3, "Ville de Trifouillis", align=GAUCHE, saisie=True)
cellule(ws_t, 7, 2, "Habitants", align=GAUCHE)
cellule(ws_t, 7, 3, 12000, fmt=FMT_ENTIER, align=DROITE, saisie=True)
DEFAUT_TB = {"carte": "Oui", "travaux": "Oui", "chantiers": "Non", "participer": "Non", "diagnostic": "Non"}
for i, (cle, nom, _, _) in enumerate(MODULES):
    cellule(ws_t, 8 + i, 2, nom, align=GAUCHE)
    c = cellule(ws_t, 8 + i, 3, DEFAUT_TB[cle], align=CENTRE, saisie=True)
    valid_oui_non_t.add(c)
cellule(ws_t, 13, 2, "Engagement, en années", align=GAUCHE)
c = cellule(ws_t, 13, 3, 3, fmt="0", align=CENTRE, saisie=True)
valid_annees_t.add(c)
ANNEES_TB_LU = ANNEES_TB

cellule(ws_t, 16, 2, "Réglages modifiés par rapport au site", align=GAUCHE, font=FONT_NOTE)
cellule(ws_t, 16, 3, f'=COUNTIF({PLAGE_ECARTS},"modifié")', fmt="0", align=CENTRE, font=FONT_NOTE)
ws_t.conditional_formatting.add("C16", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="B91C1C", bold=True, size=9)))
c = ws_t.cell(row=17, column=2, value="Zéro veut dire que ce classeur applique exactement la grille en production. Les réglages sont dans la feuille du même nom.")
c.font, c.alignment = FONT_NOTE, HAUT
ws_t.merge_cells(start_row=17, start_column=2, end_row=18, end_column=3)


def tuile(ws, ligne, colonne, libelle, formule, fmt, note=None, forte=False, largeur=1):
    fill = FILL_TUILE_FORTE if forte else FILL_TUILE
    for dr in range(3):
        for dc in range(largeur):
            cc = ws.cell(row=ligne + dr, column=colonne + dc)
            cc.fill = fill
            cc.border = BORDURE_TUILE
    if largeur > 1:
        for dr in range(3):
            ws.merge_cells(start_row=ligne + dr, start_column=colonne, end_row=ligne + dr, end_column=colonne + largeur - 1)
    c = ws.cell(row=ligne, column=colonne, value=libelle)
    c.font, c.alignment = (FONT_TUILE_LIBELLE_CLAIR if forte else FONT_TUILE_LIBELLE), Alignment(horizontal="left", vertical="bottom")
    c = ws.cell(row=ligne + 1, column=colonne, value=formule)
    c.font, c.number_format, c.alignment = (FONT_TUILE_VALEUR_CLAIR if forte else FONT_TUILE_VALEUR), fmt, Alignment(horizontal="left", vertical="center", shrink_to_fit=True)
    c = ws.cell(row=ligne + 2, column=colonne, value=note)
    c.font, c.alignment = (FONT_TUILE_NOTE_CLAIR if forte else FONT_NOTE), Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[ligne].height = 16
    ws.row_dimensions[ligne + 1].height = 34
    ws.row_dimensions[ligne + 2].height = 26


c = ws_t.cell(row=5, column=5, value="Le prix pour cette commune")
c.font = FONT_SECTION
tuile(ws_t, 6, 5, "ABONNEMENT MENSUEL HT", f"={A_TBX['mensuel']}", FMT_EUROS,
      f'=TEXT({A_TBX["n"]},"0")&IF({A_TBX["n"]}>1," modules, "," module, ")&TEXT({ANNEES_TB_LU},"0")&IF({ANNEES_TB_LU}>1," ans"," an")', forte=True)
tuile(ws_t, 6, 6, "ABONNEMENT ANNUEL HT", f"={A_TBX['annuel']}", FMT_EUROS, "Chaque année de l'engagement")
tuile(ws_t, 6, 7, "MISE EN SERVICE", f"={A_TBX['setup']}", FMT_EUROS, "Une seule fois, la première année")
tuile(ws_t, 6, 8, "TOTAL SUR LA DURÉE", f"={A_TBX['total']}", FMT_EUROS, "Mise en service comprise, hors taxes")
tuile(ws_t, 10, 5, "CE QUE LE VISITEUR VOIT SUR LE SITE", f'="de "&TRIM(TEXT({A_TBX["bas_m"]},"# ##0"))&" à "&TRIM(TEXT({A_TBX["haut_m"]},"# ##0"))&" € par mois"', "@",
      "La fourchette publique, le tarif exact se demande", largeur=2)
tuile(ws_t, 10, 7, "COMMANDE PUBLIQUE", f"={A_TBX['seuil']}", "@",
      f'=IF({A_TBX["seuil"]}="Sous les seuils","Commande sans mise en concurrence possible","Le total dépasse "&TRIM(TEXT(IF({A_TBX["total"]}>=Seuil_3,Seuil_3,IF({A_TBX["total"]}>=Seuil_2,Seuil_2,Seuil_1)),"# ##0"))&" € HT")', largeur=2)
ws_t.conditional_formatting.add("G11", CellIsRule(operator="equal", formula=['"Sous les seuils"'], font=Font(color="15803D", bold=True, size=22)))
ws_t.conditional_formatting.add("G11", CellIsRule(operator="notEqual", formula=['"Sous les seuils"'], font=Font(color="B45309", bold=True, size=22)))

# Le détail, module par module
ws_t.cell(row=14, column=5, value="Le détail du prix, par mois").font = FONT_SECTION
entete(ws_t, 15, ["Poste", "Montant", "", ""], debut=5)
ws_t.merge_cells(start_row=15, start_column=6, end_row=15, end_column=8)
r = 16
for cle, nom, _, _ in MODULES:
    cellule(ws_t, r, 5, nom, align=GAUCHE)
    cellule(ws_t, r, 6, f'=IF({A_TBX[f"prix_{cle}"]}=0,"non retenu",{A_TBX[f"prix_{cle}"]})', fmt=FMT_EUROS, align=DROITE)
    ws_t.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
    r += 1
lignes_detail_tb = [
    ("Remise multi-modules, sur le plus cher", f'=IF({A_TBX["rem_mod"]}=0,"aucune",-{A_TBX["rem_mod"]})', FONT_NORMAL),
    ("Remise d'engagement", f'=IF({A_TBX["rem_eng"]}=0,"aucune",-{A_TBX["rem_eng"]})', FONT_NORMAL),
    ("Abonnement mensuel HT", f"={A_TBX['mensuel']}", FONT_GRAS),
]
for libelle, formule, font in lignes_detail_tb:
    cellule(ws_t, r, 5, libelle, align=GAUCHE, font=font)
    cellule(ws_t, r, 6, formule, fmt=FMT_EUROS_SIGNE if "Remise" in libelle else FMT_EUROS, align=DROITE, font=font)
    ws_t.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
    r += 1
for rr in range(16, r):
    ws_t.column_dimensions["E"].width = 34

# Le graphique 1 : le prix de chaque module dans ce scénario
g1 = BarChart()
g1.type = "bar"
g1.style = 10
g1.title = "Le prix mensuel de chaque module retenu"
g1.y_axis.title = None
g1.x_axis.title = None
g1.y_axis.number_format = '#,##0" €"'
g1.legend = None
g1.height, g1.width = 7.5, 13.5
donnees = Reference(ws_c, min_col=2, min_row=19, max_row=19 + len(MODULES))
cats = Reference(ws_c, min_col=1, min_row=20, max_row=19 + len(MODULES))
g1.add_data(donnees, titles_from_data=True)
g1.set_categories(cats)
g1.series[0].graphicalProperties.solidFill = ROUGE
g1.series[0].graphicalProperties.line.noFill = True
g1.dataLabels = DataLabelList()
g1.dataLabels.showVal = True
g1.dataLabels.numFmt = '#,##0" €"'
g1.gapWidth = 60
ws_t.add_chart(g1, "J5")

# Le graphique 2 : le même scénario sur les quatre durées d'engagement
g2 = BarChart()
g2.type = "col"
g2.style = 10
g2.title = "Ce que change la durée d'engagement"
g2.y_axis.number_format = '#,##0" €"'
g2.height, g2.width = 7.5, 13.5
donnees = Reference(ws_c, min_col=3, min_row=12, max_row=16)
cats = Reference(ws_c, min_col=1, min_row=13, max_row=16)
g2.add_data(donnees, titles_from_data=True)
g2.set_categories(cats)
g2.series[0].graphicalProperties.solidFill = ENCRE
g2.series[0].graphicalProperties.line.noFill = True
g2.dataLabels = DataLabelList()
g2.dataLabels.showVal = True
g2.dataLabels.numFmt = '#,##0" €"'
g2.legend.position = "b"
g2.gapWidth = 80
ws_t.add_chart(g2, "J21")

# Le graphique 3 : la courbe de la suite complète et des modules seuls selon la taille
g3 = LineChart()
g3.style = 12
g3.title = "L'abonnement mensuel selon la taille de la commune (grille, engagement de la feuille Grille)"
g3.y_axis.number_format = '#,##0" €"'
g3.x_axis.title = "Habitants"
g3.height, g3.width = 9, 27.5
for j, (cle, nom, _, couleur) in enumerate(MODULES):
    ref = Reference(ws_g, min_col=4 + j, min_row=7, max_row=LIGNE_G1)
    g3.add_data(ref, titles_from_data=True)
    g3.series[-1].graphicalProperties.line.solidFill = couleur
    g3.series[-1].graphicalProperties.line.width = 22000
    g3.series[-1].smooth = True
ref = Reference(ws_g, min_col=COL_SUITE_G, min_row=7, max_row=LIGNE_G1)
g3.add_data(ref, titles_from_data=True)
g3.series[-1].graphicalProperties.line.solidFill = ROUGE
g3.series[-1].graphicalProperties.line.width = 34000
g3.series[-1].smooth = True
g3.set_categories(Reference(ws_g, min_col=2, min_row=LIGNE_G0, max_row=LIGNE_G1))
g3.legend.position = "b"
ws_t.add_chart(g3, "B27")

# Le graphique 4 : les scénarios de la feuille Scénarios, totaux sur la durée
g4 = BarChart()
g4.type = "bar"
g4.style = 10
g4.title = "Vos scénarios : le total sur la durée, mise en service comprise"
g4.y_axis.number_format = '#,##0" €"'
g4.legend = None
g4.height, g4.width = 11, 27.5
donnees = Reference(ws_s, min_col=COL_SORTIE_S + 3, min_row=5, max_row=LIGNE_S1)
cats = Reference(ws_s, min_col=2, min_row=LIGNE_S0, max_row=LIGNE_S1)
g4.add_data(donnees, titles_from_data=True)
g4.set_categories(cats)
g4.series[0].graphicalProperties.solidFill = "1B5FA8"
g4.series[0].graphicalProperties.line.noFill = True
g4.dataLabels = DataLabelList()
g4.dataLabels.showVal = True
g4.dataLabels.numFmt = '#,##0" €"'
g4.gapWidth = 40
ws_t.add_chart(g4, "B46")

c = ws_t.cell(row=68, column=2, value=(
    "Les scénarios vides de la feuille Scénarios apparaissent sans barre. Les seuils de la commande publique sont à 60 000, "
    "90 000 et 216 000 euros hors taxes sur toute la durée."
))
c.font, c.alignment = FONT_NOTE, GAUCHE
ws_t.merge_cells(start_row=68, start_column=2, end_row=68, end_column=8)
proteger(ws_t)

# ---------------------------------------------------------------------------
# Feuille Lisez-moi
# ---------------------------------------------------------------------------

ws_l = wb.create_sheet("Lisez-moi", 1)
ws_l.sheet_view.showGridLines = False
largeurs(ws_l, [3, 110])
titre(ws_l, "Comment se servir de ce classeur")
PARAGRAPHES = [
    "Ce classeur reproduit le calcul de prix de la page de tarification du site, tel qu'il est en production. Il sert "
    "à préparer un rendez-vous, à comparer des offres et à voir ce que donnerait un changement de grille. Rien n'est "
    "à renvoyer : c'est votre outil de travail.",
    "Les cases jaunes se remplissent, tout le reste se calcule. La feuille Tableau de bord affiche le prix d'une "
    "commune : sa population, les modules retenus, la durée d'engagement, puis l'abonnement, la mise en service, le "
    "total sur la durée, la fourchette que le visiteur voit sur le site et la situation face aux seuils de la commande "
    "publique. Les graphiques suivent le scénario.",
    "La feuille Scénarios accepte jusqu'à vingt-cinq cas, un par ligne, avec le prix de chacun et sa fourchette "
    "publique. Le dernier graphique du tableau de bord les compare.",
    "La feuille Grille montre le prix sur toutes les tailles de communes, module par module puis pour la suite "
    "complète. C'est elle qui nourrit la courbe du tableau de bord.",
    "La feuille Réglages contient la grille de production et la valeur que ce classeur utilise. Vous pouvez modifier "
    "cette dernière pour voir ce qu'un autre prix, un autre poids ou une autre remise donnerait. Le tableau de bord "
    "compte les réglages modifiés : à zéro, vous êtes exactement sur la grille du site. Pour revenir en arrière, "
    "recopiez la valeur grise.",
    "Les feuilles sont protégées sans mot de passe, pour éviter d'effacer une formule par mégarde. L'onglet Révision "
    "d'Excel permet d'ôter la protection si vous voulez aller plus loin. Les colonnes de calcul détaillé sont repliées "
    "à droite des feuilles Grille et Scénarios.",
    "Tous les montants sont hors taxes. La TVA de 20 % s'ajoute.",
]
ligne = 5
for texte in PARAGRAPHES:
    c = ws_l.cell(row=ligne, column=2, value=texte)
    c.font, c.alignment = FONT_NORMAL, HAUT
    ws_l.row_dimensions[ligne].height = 15 * (len(texte) // 105 + 1) + 6
    ligne += 2
c = ws_l.cell(row=ligne, column=2, value="Case jaune : à remplir par vous")
c.fill, c.font = FILL_SAISIE, FONT_NORMAL
ligne += 1
c = ws_l.cell(row=ligne, column=2, value="Case grise : la valeur en production sur le site, en lecture seule")
c.fill, c.font = FILL_PROD, FONT_NORMAL

# Une feuille par défaut inutile
del wb["Sheet"]
wb.active = 0

# ---------------------------------------------------------------------------
# Enregistrement
# ---------------------------------------------------------------------------

wb.calculation.fullCalcOnLoad = True
sortie = Path(__file__).with_name("simulateur-tarification.xlsx")
wb.save(sortie)
print(f"Classeur écrit : {sortie}")
