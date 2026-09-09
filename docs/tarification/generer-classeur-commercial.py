#!/usr/bin/env python3
"""Génère le classeur Excel que l'équipe commerciale utilise pour proposer une
nouvelle grille de prix.

Le classeur reproduit, en formules Excel, le modèle de prix de
`home-src/src/v2/data/tarification.js` : même courbe en puissance, mêmes
poids, mêmes remises, même mise en service. Chaque réglage y figure deux
fois : la valeur du site aujourd'hui (grisée, à ne pas toucher) et la
proposition (jaune, à remplir). Tous les calculs lisent la proposition, et les
écarts avec le site s'affichent à côté.

Quand la grille du site change, on relance ce script pour remettre la colonne
« site aujourd'hui » à jour :

    python3 docs/tarification/generer-classeur-commercial.py

Le classeur sort à côté du script : `grille-de-prix-commerciaux.xlsx`.
Dépendance : openpyxl (pip install openpyxl)."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------------------
# La grille du site aujourd'hui : la copie exacte de tarification.js.
# Ce bloc est la seule chose à mettre à jour quand la grille du site change.
# ---------------------------------------------------------------------------

ANCRE = {"population": 12000, "prix": 200}
EXPOSANT = 0.5
POPULATION = {"min": 500, "max": 2500000}
MODULES = [
    ("carte", "Carte des projets", 1),
    ("travaux", "Travaux du quotidien", 0.6),
    ("participer", "Participer", 2),
    ("diagnostic", "Diagnostic terrain", 0.8),
    ("chantiers", "Chantiers et arrêtés", 3),
]
CHANTIERS_DEMI_POIDS = {"sous": 5000, "plein": 20000}
MISE_EN_SERVICE = {"mois": 6, "offerteSous": 0}
REMISE_PAR_MODULE_AJOUTE = 0.10
ENGAGEMENTS = [(1, 0), (2, 0.10), (3, 0.15), (4, 0.20)]
SEUILS = [
    (60000, "Procédure adaptée"),
    (90000, "Publicité obligatoire"),
    (216000, "Procédure formalisée"),
]
REPERES = [
    ("Village", 800),
    ("Bourg", 3000),
    ("Petite ville", 12000),
    ("Ville moyenne", 50000),
    ("Grande ville", 150000),
    ("Métropole", 500000),
]

# ---------------------------------------------------------------------------
# Mise en forme
# ---------------------------------------------------------------------------

FILL_SAISIE = PatternFill("solid", fgColor="FFF2CC")  # jaune : à remplir
FILL_SITE = PatternFill("solid", fgColor="EDEDED")  # gris : le site aujourd'hui
FILL_ENTETE = PatternFill("solid", fgColor="1F2937")
FILL_SECTION = PatternFill("solid", fgColor="E5E7EB")
FILL_RESULTAT = PatternFill("solid", fgColor="F0FDF4")
FONT_TITRE = Font(name="Calibri", size=16, bold=True, color="1F2937")
FONT_ENTETE = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_SECTION = Font(name="Calibri", size=11, bold=True, color="1F2937")
FONT_GRAS = Font(name="Calibri", size=11, bold=True)
FONT_NOTE = Font(name="Calibri", size=10, italic=True, color="6B7280")
FONT_NORMAL = Font(name="Calibri", size=11)
BORDURE = Border(*(Side(style="thin", color="D1D5DB"),) * 4)
GAUCHE = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)
DROITE = Alignment(horizontal="right", vertical="center")
HAUT = Alignment(horizontal="left", vertical="top", wrap_text=True)

FMT_EUROS = '#,##0" €"'
FMT_EUROS_SIGNE = '+#,##0" €";-#,##0" €";0" €"'
FMT_POURCENT = "0%"
FMT_POURCENT_SIGNE = "+0%;-0%;0%"
FMT_ENTIER = "#,##0"
FMT_DECIMAL = "0.00"

DEVERROUILLE = Protection(locked=False)

SAISIE_MODULE = {"carte": "Oui", "travaux": "Oui", "participer": "Non", "diagnostic": "Non", "chantiers": "Non"}

wb = Workbook()


def nom_defini(nom, feuille, plage):
    wb.defined_names[nom] = DefinedName(nom, attr_text=f"'{feuille}'!{plage}")


def entete(ws, ligne, colonnes, debut=1):
    for i, texte in enumerate(colonnes):
        c = ws.cell(row=ligne, column=debut + i, value=texte)
        c.fill, c.font, c.alignment, c.border = FILL_ENTETE, FONT_ENTETE, CENTRE, BORDURE


def cellule(ws, ligne, colonne, valeur=None, fmt=None, fill=None, font=FONT_NORMAL, align=None, saisie=False):
    c = ws.cell(row=ligne, column=colonne, value=valeur)
    c.font, c.border = font, BORDURE
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if saisie:
        c.fill = FILL_SAISIE
        c.protection = DEVERROUILLE
    return c


def largeurs(ws, valeurs):
    for i, largeur in enumerate(valeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur


def proteger(ws):
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.sort = False
    ws.protection.autoFilter = False


# ---------------------------------------------------------------------------
# Le calcul, écrit une fois, posé sur n'importe quelle feuille.
#
# `N` vaut "" pour la proposition et "Site_" pour la grille du site : c'est le
# préfixe des noms définis sur la feuille Réglages. `a` donne l'adresse des
# cellules intermédiaires ; `sel` l'adresse de la case Oui / Non de chaque
# module (None = module toujours retenu) ; `n` l'expression qui compte les
# modules retenus ; `annees` l'adresse ou la valeur de l'engagement.
# ---------------------------------------------------------------------------

def formules(N, a, sel, n, annees, population):
    eng_col = 2 if N else 3  # colonne de la table des engagements : site ou proposition
    f = {}
    f["pb"] = f"=MIN({N}Population_maximum,MAX({N}Population_minimum,ROUND({population},0)))"
    f["unite"] = f"={N}Prix_ancre*({a['pb']}/{N}Population_ancre)^{N}Exposant"
    f["pch"] = (
        f"=IF({a['pb']}<={N}Chantiers_demi_poids_sous,{N}Poids_chantiers/2,"
        f"IF({a['pb']}>={N}Chantiers_plein_poids_des,{N}Poids_chantiers,"
        f"{N}Poids_chantiers*(0.5+0.5*(LN({a['pb']})-LN({N}Chantiers_demi_poids_sous))"
        f"/(LN({N}Chantiers_plein_poids_des)-LN({N}Chantiers_demi_poids_sous)))))"
    )
    for cle, _, _ in MODULES:
        poids = a["pch"] if cle == "chantiers" else f"{N}Poids_{cle}"
        prix = f"{a['unite']}*{poids}"
        f[f"prix_{cle}"] = f"=IF({sel[cle]}=\"Oui\",{prix},0)" if sel.get(cle) else f"={prix}"
    f["brut"] = f"=SUM({a['prix_range']})"
    f["n"] = f"={n}"
    f["taux_mod"] = f"=MIN(0.9,MAX(0,{a['n']}-1)*{N}Remise_par_module_ajoute)"
    f["rem_mod"] = f"=MAX({a['prix_range']})*{a['taux_mod']}"
    f["apres"] = f"={a['brut']}-{a['rem_mod']}"
    f["taux_eng"] = f"=IFERROR(VLOOKUP({annees},Table_engagements,{eng_col},FALSE),0)"
    f["rem_eng"] = f"={a['apres']}*{a['taux_eng']}"
    f["mensuel"] = f"={a['apres']}-{a['rem_eng']}"
    f["annuel"] = f"={a['mensuel']}*12"
    f["setup"] = f"=IF({a['pb']}<{N}Mise_en_service_offerte_sous,0,{a['mensuel']}*{N}Mise_en_service_mois)"
    f["total"] = f"={a['setup']}+{a['annuel']}*IF(N({annees})>0,{annees},1)"
    f["seuil"] = (
        f"=IF({a['total']}>=Seuil_3,\"{SEUILS[2][1]}\","
        f"IF({a['total']}>=Seuil_2,\"{SEUILS[1][1]}\","
        f"IF({a['total']}>=Seuil_1,\"{SEUILS[0][1]}\",\"Sous les seuils\")))"
    )
    return f


ORDRE_INTERMEDIAIRES = ["pb", "unite", "pch"] + [f"prix_{c}" for c, _, _ in MODULES] + [
    "brut", "n", "taux_mod", "rem_mod", "apres", "taux_eng", "rem_eng", "mensuel", "annuel", "setup", "total", "seuil",
]
LIBELLES_INTERMEDIAIRES = {
    "pb": "Population retenue", "unite": "Prix d'une unité de poids", "pch": "Poids chantiers",
    "brut": "Total avant remises", "n": "Modules retenus", "taux_mod": "Taux de remise multi-modules",
    "rem_mod": "Remise multi-modules", "apres": "Après remise multi-modules", "taux_eng": "Taux engagement",
    "rem_eng": "Remise engagement", "mensuel": "Mensuel", "annuel": "Annuel", "setup": "Mise en service",
    "total": "Total sur la durée", "seuil": "Seuil",
}
for cle, nom, _ in MODULES:
    LIBELLES_INTERMEDIAIRES[f"prix_{cle}"] = f"Prix {nom}"
FORMATS_INTERMEDIAIRES = {
    "pb": FMT_ENTIER, "unite": FMT_EUROS, "pch": FMT_DECIMAL, "brut": FMT_EUROS, "n": "0", "taux_mod": FMT_POURCENT,
    "rem_mod": FMT_EUROS, "apres": FMT_EUROS, "taux_eng": FMT_POURCENT, "rem_eng": FMT_EUROS, "mensuel": FMT_EUROS,
    "annuel": FMT_EUROS, "setup": FMT_EUROS, "total": FMT_EUROS, "seuil": "@",
}
for cle, _, _ in MODULES:
    FORMATS_INTERMEDIAIRES[f"prix_{cle}"] = FMT_EUROS


def bloc_horizontal(ws, ligne, colonne_debut, N, sel, n, annees, population, masque=True):
    """Pose le calcul sur une ligne, une colonne par intermédiaire, à partir de
    `colonne_debut`. Rend l'adresse de chaque intermédiaire."""
    a = {}
    for i, cle in enumerate(ORDRE_INTERMEDIAIRES):
        a[cle] = f"{get_column_letter(colonne_debut + i)}{ligne}"
    premier = ORDRE_INTERMEDIAIRES.index("prix_carte")
    a["prix_range"] = (
        f"{get_column_letter(colonne_debut + premier)}{ligne}:"
        f"{get_column_letter(colonne_debut + premier + len(MODULES) - 1)}{ligne}"
    )
    f = formules(N, a, sel, n, annees, population)
    for i, cle in enumerate(ORDRE_INTERMEDIAIRES):
        c = ws.cell(row=ligne, column=colonne_debut + i, value=f[cle])
        c.number_format = FORMATS_INTERMEDIAIRES[cle]
        c.font = FONT_NOTE
        if masque:
            ws.column_dimensions[get_column_letter(colonne_debut + i)].hidden = True
            ws.column_dimensions[get_column_letter(colonne_debut + i)].outlineLevel = 1
    return a


def entetes_intermediaires(ws, ligne, colonne_debut, prefixe):
    for i, cle in enumerate(ORDRE_INTERMEDIAIRES):
        c = ws.cell(row=ligne, column=colonne_debut + i, value=f"{prefixe}{LIBELLES_INTERMEDIAIRES[cle]}")
        c.font, c.alignment = FONT_NOTE, CENTRE
        ws.column_dimensions[get_column_letter(colonne_debut + i)].width = 14


# ---------------------------------------------------------------------------
# Feuille 1 : Lisez-moi
# ---------------------------------------------------------------------------

ws = wb.active
ws.title = "Lisez-moi"
ws.sheet_view.showGridLines = False
largeurs(ws, [3, 110])
ws.cell(row=2, column=2, value="Grille de prix Open Projets : le classeur de l'équipe commerciale").font = FONT_TITRE

PARAGRAPHES = [
    "Ce classeur reproduit, à l'identique, le calcul de prix de la page de tarification du site. Il sert à proposer "
    "une nouvelle grille : vous changez les réglages, vous regardez ce que cela donne sur des communes concrètes, et "
    "quand le résultat vous convient, vous nous renvoyez le fichier. Nous reportons alors vos valeurs sur le site.",
    "La règle de couleur est la même partout. Les cases jaunes sont à remplir. Les cases grises montrent ce que le site "
    "applique aujourd'hui et ne se modifient pas. Tout le reste est calculé et se met à jour tout seul.",
    "La feuille Réglages est celle qui compte : chaque ligne est un chiffre de la politique tarifaire, avec la valeur du "
    "site à gauche et votre proposition à droite. Au départ, la proposition est égale au site. Une colonne vous permet "
    "d'écrire pourquoi vous proposez ce changement, et le bas de la feuille porte votre nom et la date.",
    "La feuille Simulateur fait la même chose que la page du site : une population, les modules retenus, une durée "
    "d'engagement, et le détail du prix ligne par ligne, avec votre proposition et le site côte à côte.",
    "La feuille Grille montre votre proposition sur une quinzaine de tailles de communes à la fois : chaque module "
    "seul, puis la suite complète, avec l'écart par rapport au site. C'est elle qui dit si la courbe est cohérente "
    "d'un bout à l'autre.",
    "La feuille Scénarios vous laisse écrire jusqu'à vingt cas concrets, par exemple vos prospects en cours, et compare "
    "pour chacun votre proposition au site. Vous pouvez remplacer les exemples déjà présents.",
    "Comment le prix se forme aujourd'hui. Le prix d'un module suit la population en puissance : une commune dix "
    f"fois plus peuplée paie environ {10 ** EXPOSANT:.1f} fois plus. Toute la courbe est fixée par un seul point, le "
    f"prix d'une unité de poids pour une commune de {ANCRE['population']:,} habitants. Chaque module a un poids qui "
    f"multiplie ce prix. Le module chantiers vaut la moitié de son poids sous {CHANTIERS_DEMI_POIDS['sous']:,} "
    f"habitants et son poids entier à partir de {CHANTIERS_DEMI_POIDS['plein']:,}. Prendre plusieurs modules fait "
    f"baisser le plus cher d'entre eux de {REMISE_PAR_MODULE_AJOUTE:.0%} par module ajouté. L'engagement fait ensuite "
    f"baisser tout l'abonnement, chaque année. La mise en service coûte {MISE_EN_SERVICE['mois']} mois d'abonnement, "
    "une seule fois"
    + (f", et elle est offerte sous {MISE_EN_SERVICE['offerteSous']:,} habitants." if MISE_EN_SERVICE["offerteSous"] else ".")
    + " Tous les montants sont hors taxes.",
    "Les seuils de la commande publique (60 000, 90 000 et 216 000 euros hors taxes sur toute la durée) sont "
    "rappelés à titre d'information : ils ne se négocient pas, mais ils pèsent sur ce qu'une collectivité peut "
    "commander sans mise en concurrence. Le simulateur dit lequel est dépassé.",
    "Les feuilles sont protégées sans mot de passe pour éviter d'écraser une formule par mégarde. Si vous avez besoin "
    "d'y toucher, l'onglet Révision d'Excel permet d'ôter la protection. Les colonnes de calcul intermédiaire des "
    "feuilles Grille et Scénarios sont repliées à droite : le bouton « + » au-dessus des colonnes les déplie.",
    "Pour nous renvoyer votre proposition, il suffit d'enregistrer ce fichier et de nous l'envoyer. La feuille "
    "Réglages, avec vos commentaires, est ce que nous lisons en premier.",
]
ligne = 4
for texte in PARAGRAPHES:
    c = ws.cell(row=ligne, column=2, value=texte)
    c.font, c.alignment = FONT_NORMAL, HAUT
    ws.row_dimensions[ligne].height = 15 * (len(texte) // 105 + 1) + 6
    ligne += 2

ws.cell(row=ligne, column=2, value="Légende").font = FONT_SECTION
ligne += 1
c = ws.cell(row=ligne, column=2, value="Case jaune : à remplir par vous")
c.fill, c.font = FILL_SAISIE, FONT_NORMAL
ligne += 1
c = ws.cell(row=ligne, column=2, value="Case grise : la valeur appliquée sur le site aujourd'hui, en lecture seule")
c.fill, c.font = FILL_SITE, FONT_NORMAL
ligne += 1
c = ws.cell(row=ligne, column=2, value="Case verte : un résultat calculé à partir de votre proposition")
c.fill, c.font = FILL_RESULTAT, FONT_NORMAL

# ---------------------------------------------------------------------------
# Feuille 2 : Réglages
# ---------------------------------------------------------------------------

REGLAGES = "Réglages"
ws = wb.create_sheet(REGLAGES)
ws.sheet_view.showGridLines = False
largeurs(ws, [46, 14, 20, 20, 60])
ws.cell(row=1, column=1, value="Réglages de la grille de prix").font = FONT_TITRE
c = ws.cell(row=2, column=1, value=(
    "Chaque ligne est un chiffre de la politique tarifaire. La colonne grise montre le site aujourd'hui, la colonne "
    "jaune porte votre proposition : c'est elle que lisent toutes les autres feuilles. Écrivez en face pourquoi vous "
    "proposez ce changement."
))
c.font, c.alignment = FONT_NOTE, GAUCHE
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)
ws.row_dimensions[2].height = 34

entete(ws, 4, ["Réglage", "Précision", "Sur le site aujourd'hui", "Votre proposition", "Pourquoi ce changement"])
ws.freeze_panes = "A5"
ligne = 5


def section(titre, explication=None):
    global ligne
    c = ws.cell(row=ligne, column=1, value=titre)
    c.font, c.fill = FONT_SECTION, FILL_SECTION
    for col in range(2, 6):
        ws.cell(row=ligne, column=col).fill = FILL_SECTION
    ligne += 1
    if explication:
        c = ws.cell(row=ligne, column=1, value=explication)
        c.font, c.alignment = FONT_NOTE, GAUCHE
        ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=5)
        ws.row_dimensions[ligne].height = 15 * (len(explication) // 150 + 1) + 4
        ligne += 1


def reglage(libelle, nom, valeur, fmt, precision=None, precision_fmt=None, modifiable=True, commentaire=None):
    """Une ligne de réglage : la valeur du site en C, la proposition en D, et
    deux noms définis (`Site_<nom>` et `<nom>`) pour les formules."""
    global ligne
    cellule(ws, ligne, 1, libelle, align=GAUCHE)
    if precision is not None:
        cellule(ws, ligne, 2, precision, fmt=precision_fmt, align=CENTRE)
    else:
        cellule(ws, ligne, 2, None)
    c_site = cellule(ws, ligne, 3, valeur, fmt=fmt, fill=FILL_SITE, align=DROITE)
    c_prop = cellule(ws, ligne, 4, valeur, fmt=fmt, align=DROITE, saisie=modifiable)
    if not modifiable:
        c_prop.fill = FILL_SITE
    cellule(ws, ligne, 5, None, saisie=True, align=GAUCHE)
    if commentaire:
        c_site.comment = Comment(commentaire, "Open Projets")
    nom_defini(f"Site_{nom}", REGLAGES, f"$C${ligne}")
    nom_defini(nom, REGLAGES, f"$D${ligne}")
    ligne += 1
    return ligne - 1


section(
    "La courbe",
    "Le prix d'une unité de poids suit la population en puissance. Ces deux chiffres déplacent toute la grille : le "
    "prix à l'ancre fait monter ou baisser tous les prix en proportion, l'exposant dit à quelle vitesse le prix monte "
    "avec la population (0,6 : dix fois plus d'habitants, quatre fois plus cher ; 1 : dix fois plus cher).",
)
reglage("Population de l'ancre", "Population_ancre", ANCRE["population"], FMT_ENTIER, "habitants")
reglage("Prix mensuel HT d'une unité de poids à l'ancre", "Prix_ancre", ANCRE["prix"], FMT_EUROS, "€ HT / mois")
reglage("Exposant de la courbe", "Exposant", EXPOSANT, FMT_DECIMAL)
reglage("Population minimale du simulateur", "Population_minimum", POPULATION["min"], FMT_ENTIER, "habitants")
reglage("Population maximale du simulateur", "Population_maximum", POPULATION["max"], FMT_ENTIER, "habitants")
ligne += 1

section(
    "Le poids de chaque module",
    "Le prix d'un module est le prix de l'unité multiplié par son poids. Un poids de 1 vaut exactement le prix de "
    "l'unité, un poids de 2 le double.",
)
for cle, nom, poids in MODULES:
    reglage(nom, f"Poids_{cle}", poids, FMT_DECIMAL, "fois l'unité")
ligne += 1

section(
    "Le module chantiers pour les petites communes",
    "Le module chantiers vaut la moitié de son poids jusqu'à la première borne, son poids entier à partir de la "
    "seconde, et monte en pente douce entre les deux.",
)
reglage("Demi-poids jusqu'à", "Chantiers_demi_poids_sous", CHANTIERS_DEMI_POIDS["sous"], FMT_ENTIER, "habitants")
reglage("Poids entier à partir de", "Chantiers_plein_poids_des", CHANTIERS_DEMI_POIDS["plein"], FMT_ENTIER, "habitants")
ligne += 1

section(
    "La remise quand on prend plusieurs modules",
    "La remise porte sur le plus cher des modules retenus, et vaut ce taux par module ajouté au premier : avec deux "
    "modules le plus cher baisse de 10 %, avec trois de 20 %, avec cinq de 40 %.",
)
reglage("Remise par module ajouté, sur le module le plus cher", "Remise_par_module_ajoute", REMISE_PAR_MODULE_AJOUTE, FMT_POURCENT)
ligne += 1

section(
    "La remise d'engagement",
    "Elle s'applique à tout l'abonnement, chaque année, selon la durée d'engagement. Un accord-cadre public ne "
    "dépasse pas quatre ans.",
)
debut_eng = ligne
for annees, remise in ENGAGEMENTS:
    reglage(f"Engagement {annees} an{'s' if annees > 1 else ''}", f"Remise_engagement_{annees}_ans", remise, FMT_POURCENT, annees, "0")
fin_eng = ligne - 1
nom_defini("Table_engagements", REGLAGES, f"$B${debut_eng}:$D${fin_eng}")
ligne += 1

section(
    "La mise en service",
    "Facturée une seule fois, en mois d'abonnement tel qu'il sort après toutes les remises. Offerte aux communes sous "
    "le seuil d'habitants.",
)
reglage("Mise en service, en mois d'abonnement", "Mise_en_service_mois", MISE_EN_SERVICE["mois"], "0", "mois")
reglage("Mise en service offerte sous", "Mise_en_service_offerte_sous", MISE_EN_SERVICE["offerteSous"], FMT_ENTIER, "habitants")
ligne += 1

section(
    "Les seuils de la commande publique (pour information)",
    "Ces montants sont fixés par le code de la commande publique et s'apprécient sur le total hors taxes de tout le "
    "marché, mise en service comprise. Ils ne se modifient pas ici.",
)
for i, (montant, nom) in enumerate(SEUILS, start=1):
    reglage(nom, f"Seuil_{i}", montant, FMT_EUROS, "€ HT, total", modifiable=False)
ligne += 2

c = ws.cell(row=ligne, column=1, value="Votre proposition est portée par")
c.font = FONT_SECTION
ligne += 1
cellule(ws, ligne, 1, "Nom", align=GAUCHE)
cellule(ws, ligne, 2, None, saisie=True)
ws.merge_cells(start_row=ligne, start_column=2, end_row=ligne, end_column=4)
ligne += 1
cellule(ws, ligne, 1, "Date", align=GAUCHE)
cellule(ws, ligne, 2, None, fmt="DD/MM/YYYY", saisie=True)
ws.merge_cells(start_row=ligne, start_column=2, end_row=ligne, end_column=4)
ligne += 1
cellule(ws, ligne, 1, "Remarques générales", align=HAUT)
cellule(ws, ligne, 2, None, saisie=True, align=HAUT)
ws.merge_cells(start_row=ligne, start_column=2, end_row=ligne + 3, end_column=5)
for r in range(ligne, ligne + 4):
    ws.row_dimensions[r].height = 18
proteger(ws)

# ---------------------------------------------------------------------------
# Feuille 3 : Simulateur
# ---------------------------------------------------------------------------

SIMULATEUR = "Simulateur"
ws = wb.create_sheet(SIMULATEUR)
ws.sheet_view.showGridLines = False
largeurs(ws, [44, 22, 22, 18, 4])
ws.cell(row=1, column=1, value="Simulateur : une commune, ses modules, son prix").font = FONT_TITRE
c = ws.cell(row=2, column=1, value=(
    "Remplissez les cases jaunes comme le ferait un visiteur de la page du site. Le détail du prix s'affiche en "
    "dessous, avec votre proposition et le site côte à côte."
))
c.font, c.alignment = FONT_NOTE, GAUCHE
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
ws.row_dimensions[2].height = 30

valid_oui_non = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=False)
valid_annees = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=False)
ws.add_data_validation(valid_oui_non)
ws.add_data_validation(valid_annees)

entete(ws, 4, ["La commune et son choix", "Votre saisie"])
cellule(ws, 5, 1, "Population de la commune", align=GAUCHE)
cellule(ws, 5, 2, 12000, fmt=FMT_ENTIER, align=DROITE, saisie=True)
POP = "$B$5"
sel = {}
r = 6
for cle, nom, _ in MODULES:
    cellule(ws, r, 1, f"Module {nom}", align=GAUCHE)
    c = cellule(ws, r, 2, SAISIE_MODULE[cle], align=CENTRE, saisie=True)
    valid_oui_non.add(c)
    sel[cle] = f"$B${r}"
    r += 1
cellule(ws, r, 1, "Durée d'engagement, en années", align=GAUCHE)
c = cellule(ws, r, 2, 3, fmt="0", align=CENTRE, saisie=True)
valid_annees.add(c)
ANNEES = f"$B${r}"
SEL_RANGE = f"$B$6:$B${r - 1}"
r += 2

entete(ws, r, ["Le détail du prix, hors taxes", "Votre proposition", "Le site aujourd'hui", "Écart"])
r += 1

# Les lignes du détail, dans l'ordre de la page du site
LIGNES_DETAIL = [
    ("unite", "Prix d'une unité de poids pour cette population", FMT_EUROS, FMT_EUROS_SIGNE, False),
    ("pch", "Poids appliqué au module chantiers", FMT_DECIMAL, "+0.00;-0.00;0", False),
] + [(f"prix_{cle}", f"{nom}, par mois", FMT_EUROS, FMT_EUROS_SIGNE, False) for cle, nom, _ in MODULES] + [
    ("brut", "Total des modules avant remises, par mois", FMT_EUROS, FMT_EUROS_SIGNE, True),
    ("n", "Nombre de modules retenus", "0", "+0;-0;0", False),
    ("taux_mod", "Taux de la remise multi-modules (sur le plus cher)", FMT_POURCENT, FMT_POURCENT_SIGNE, False),
    ("rem_mod", "Remise multi-modules, par mois", FMT_EUROS, FMT_EUROS_SIGNE, False),
    ("apres", "Après remise multi-modules, par mois", FMT_EUROS, FMT_EUROS_SIGNE, False),
    ("taux_eng", "Taux de la remise d'engagement", FMT_POURCENT, FMT_POURCENT_SIGNE, False),
    ("rem_eng", "Remise d'engagement, par mois", FMT_EUROS, FMT_EUROS_SIGNE, False),
    ("mensuel", "Abonnement mensuel", FMT_EUROS, FMT_EUROS_SIGNE, True),
    ("annuel", "Abonnement annuel", FMT_EUROS, FMT_EUROS_SIGNE, True),
    ("setup", "Mise en service, une seule fois", FMT_EUROS, FMT_EUROS_SIGNE, False),
    ("total", "Total sur toute la durée d'engagement", FMT_EUROS, FMT_EUROS_SIGNE, True),
    ("seuil", "Seuil de la commande publique dépassé", "@", None, True),
]
adresses = {"": {}, "Site_": {}}
debut_detail = r
lignes_detail = {}
for cle, libelle, fmt, fmt_ecart, gras in LIGNES_DETAIL:
    lignes_detail[cle] = r
    adresses[""][cle] = f"$B${r}"
    adresses["Site_"][cle] = f"$C${r}"
    r += 1
# La population retenue (bornée) se calcule en dehors du tableau, sur une ligne discrète
for N, col in (("", 2), ("Site_", 3)):
    adresses[N]["pb"] = f"${get_column_letter(col)}${r}"
    adresses[N]["prix_range"] = (
        f"${get_column_letter(col)}${lignes_detail['prix_carte']}:${get_column_letter(col)}${lignes_detail['prix_chantiers']}"
    )
cellule(ws, r, 1, "Population retenue par le simulateur (entre ses bornes)", font=FONT_NOTE, align=GAUCHE)

for N, col in (("", 2), ("Site_", 3)):
    f = formules(N, adresses[N], sel, f'COUNTIF({SEL_RANGE},"Oui")', ANNEES, POP)
    for cle, libelle, fmt, fmt_ecart, gras in LIGNES_DETAIL:
        fill = FILL_RESULTAT if N == "" else FILL_SITE
        cellule(ws, lignes_detail[cle], col, f[cle], fmt=fmt, fill=fill, align=DROITE if fmt != "@" else CENTRE,
                font=FONT_GRAS if gras else FONT_NORMAL)
    c = ws.cell(row=r, column=col, value=f["pb"])
    c.number_format, c.font = FMT_ENTIER, FONT_NOTE

for cle, libelle, fmt, fmt_ecart, gras in LIGNES_DETAIL:
    rr = lignes_detail[cle]
    cellule(ws, rr, 1, libelle, align=GAUCHE, font=FONT_GRAS if gras else FONT_NORMAL)
    if fmt_ecart:
        cellule(ws, rr, 4, f"=B{rr}-C{rr}", fmt=fmt_ecart, align=DROITE, font=FONT_GRAS if gras else FONT_NORMAL)
    else:
        cellule(ws, rr, 4, f'=IF(B{rr}=C{rr},"","change")', align=CENTRE, font=FONT_NOTE)

r += 2
c = ws.cell(row=r, column=1, value=(
    "L'écart est votre proposition moins le site : un montant positif veut dire que la commune paierait plus qu'aujourd'hui."
))
c.font, c.alignment = FONT_NOTE, GAUCHE
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
ws.freeze_panes = "A5"
proteger(ws)

# ---------------------------------------------------------------------------
# Feuille 4 : Grille
# ---------------------------------------------------------------------------

GRILLE = "Grille"
ws = wb.create_sheet(GRILLE)
ws.sheet_view.showGridLines = False
ws.cell(row=1, column=1, value="Grille : votre proposition sur toutes les tailles de communes").font = FONT_TITRE
c = ws.cell(row=2, column=1, value=(
    "Chaque ligne est une taille de commune (les populations et leurs noms se modifient). Les prix sont mensuels hors "
    "taxes : chaque module pris seul, puis la suite complète des cinq modules avec sa remise multi-modules. La durée "
    "d'engagement choisie ci-dessous s'applique à toute la feuille."
))
c.font, c.alignment = FONT_NOTE, GAUCHE
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=12)
ws.row_dimensions[2].height = 34
cellule(ws, 3, 1, "Durée d'engagement appliquée à la grille, en années", align=GAUCHE)
valid_annees_grille = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=False)
ws.add_data_validation(valid_annees_grille)
c = cellule(ws, 3, 2, 1, fmt="0", align=CENTRE, saisie=True)
valid_annees_grille.add(c)
ANNEES_GRILLE = "$B$3"

colonnes_grille = ["Commune type", "Habitants", "Une unité de poids"] + [f"{nom} seul" for _, nom, _ in MODULES] + [
    "Les 5 modules, par mois", "Mise en service (5 modules)", "Total sur la durée (5 modules)", "Seuil dépassé",
    "Les 5 modules sur le site", "Écart par mois", "Écart en %",
]
entete(ws, 5, colonnes_grille)
ws.row_dimensions[5].height = 46
largeurs(ws, [22, 12, 14] + [16] * len(MODULES) + [16, 16, 18, 20, 16, 14, 10])
ws.freeze_panes = "C6"

POPULATIONS_GRILLE = [
    ("Hameau", 500), ("Village", 800), ("Village", 1500), ("Bourg", 3000), ("Bourg", 5000), ("Petite ville", 8000),
    ("Petite ville", 12000), ("Petite ville", 20000), ("Ville moyenne", 35000), ("Ville moyenne", 50000),
    ("Ville moyenne", 80000), ("Grande ville", 150000), ("Grande ville", 300000), ("Métropole", 500000),
    ("Métropole", 1000000), ("Métropole", 2500000),
]
COL_INTER_P = 20  # colonne T : début des intermédiaires de la proposition
COL_INTER_S = COL_INTER_P + len(ORDRE_INTERMEDIAIRES) + 1
entetes_intermediaires(ws, 5, COL_INTER_P, "Proposition : ")
entetes_intermediaires(ws, 5, COL_INTER_S, "Site : ")
ws.column_dimensions[get_column_letter(COL_INTER_P - 1)].width = 3

r = 6
for nom, pop in POPULATIONS_GRILLE:
    cellule(ws, r, 1, nom, align=GAUCHE, saisie=True)
    cellule(ws, r, 2, pop, fmt=FMT_ENTIER, align=DROITE, saisie=True)
    pop_ref = f"$B{r}"
    aP = bloc_horizontal(ws, r, COL_INTER_P, "", {}, len(MODULES), ANNEES_GRILLE, pop_ref)
    aS = bloc_horizontal(ws, r, COL_INTER_S, "Site_", {}, len(MODULES), ANNEES_GRILLE, pop_ref)
    cellule(ws, r, 3, f"={aP['unite']}", fmt=FMT_EUROS, align=DROITE)
    for i, (cle, _, _) in enumerate(MODULES):
        # Un module seul : son prix, avec la seule remise d'engagement
        cellule(ws, r, 4 + i, f"={aP[f'prix_{cle}']}*(1-{aP['taux_eng']})", fmt=FMT_EUROS, align=DROITE)
    col = 4 + len(MODULES)
    cellule(ws, r, col, f"={aP['mensuel']}", fmt=FMT_EUROS, fill=FILL_RESULTAT, align=DROITE, font=FONT_GRAS)
    cellule(ws, r, col + 1, f"={aP['setup']}", fmt=FMT_EUROS, fill=FILL_RESULTAT, align=DROITE)
    cellule(ws, r, col + 2, f"={aP['total']}", fmt=FMT_EUROS, fill=FILL_RESULTAT, align=DROITE)
    cellule(ws, r, col + 3, f"={aP['seuil']}", fill=FILL_RESULTAT, align=CENTRE, font=FONT_NOTE)
    cellule(ws, r, col + 4, f"={aS['mensuel']}", fmt=FMT_EUROS, fill=FILL_SITE, align=DROITE)
    cellule(ws, r, col + 5, f"={aP['mensuel']}-{aS['mensuel']}", fmt=FMT_EUROS_SIGNE, align=DROITE)
    cellule(ws, r, col + 6, f"=IF({aS['mensuel']}=0,0,{aP['mensuel']}/{aS['mensuel']}-1)", fmt=FMT_POURCENT_SIGNE, align=DROITE)
    r += 1

# L'écart se colore : vert quand la proposition baisse le prix, rouge quand elle le monte
plage_ecart = f"{get_column_letter(4 + len(MODULES) + 5)}6:{get_column_letter(4 + len(MODULES) + 6)}{r - 1}"
ws.conditional_formatting.add(plage_ecart, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="B91C1C")))
ws.conditional_formatting.add(plage_ecart, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="15803D")))
r += 1
c = ws.cell(row=r, column=1, value=(
    "Les colonnes de calcul détaillé sont repliées à droite de la feuille : le bouton « + » au-dessus des colonnes les "
    "déplie si vous voulez vérifier un chiffre."
))
c.font, c.alignment = FONT_NOTE, GAUCHE
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
ws.sheet_properties.outlinePr.summaryRight = False
proteger(ws)

# ---------------------------------------------------------------------------
# Feuille 5 : Scénarios
# ---------------------------------------------------------------------------

SCENARIOS = "Scénarios"
ws = wb.create_sheet(SCENARIOS)
ws.sheet_view.showGridLines = False
ws.cell(row=1, column=1, value="Scénarios : vos cas concrets, proposition et site côte à côte").font = FONT_TITRE
c = ws.cell(row=2, column=1, value=(
    "Une ligne par cas : le nom de la commune ou du prospect, sa population, Oui ou Non pour chaque module, la durée "
    "d'engagement. Le prix se calcule avec votre proposition, puis avec la grille du site, et l'écart s'affiche. "
    "Remplacez les exemples par vos dossiers en cours."
))
c.font, c.alignment = FONT_NOTE, GAUCHE
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=17)
ws.row_dimensions[2].height = 34

colonnes_sc = ["Commune ou prospect", "Habitants"] + [nom for _, nom, _ in MODULES] + ["Engagement (années)"] + [
    "Mensuel proposé", "Annuel proposé", "Mise en service proposée", "Total proposé sur la durée", "Seuil dépassé",
    "Mensuel sur le site", "Total sur le site", "Écart par mois", "Écart en %",
]
entete(ws, 4, colonnes_sc)
ws.row_dimensions[4].height = 46
largeurs(ws, [26, 11] + [13] * len(MODULES) + [12, 14, 14, 14, 16, 20, 14, 14, 13, 10])
ws.freeze_panes = "C5"

valid_oui_non_sc = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
valid_annees_sc = DataValidation(type="list", formula1='"1,2,3,4"', allow_blank=True)
ws.add_data_validation(valid_oui_non_sc)
ws.add_data_validation(valid_annees_sc)

EXEMPLES = [
    ("Village, carte seule", 800, ["carte"], 3),
    ("Bourg, carte et travaux", 3000, ["carte", "travaux"], 3),
    ("Bourg, chantiers seul", 3000, ["chantiers"], 1),
    ("Petite ville, carte, travaux, participer", 12000, ["carte", "travaux", "participer"], 3),
    ("Petite ville, la suite complète", 12000, ["carte", "travaux", "participer", "diagnostic", "chantiers"], 4),
    ("Ville moyenne, carte et participer", 50000, ["carte", "participer"], 2),
    ("Ville moyenne, la suite complète", 50000, ["carte", "travaux", "participer", "diagnostic", "chantiers"], 3),
    ("Grande ville, carte et diagnostic", 150000, ["carte", "diagnostic"], 3),
    ("Métropole, carte seule", 500000, ["carte"], 4),
]
NB_LIGNES_SC = 20
COL_ENG = 3 + len(MODULES)  # colonne H
COL_SORTIE = COL_ENG + 1  # colonne I
COL_INTER_P = COL_SORTIE + 9 + 1  # une colonne vide après les sorties
COL_INTER_S = COL_INTER_P + len(ORDRE_INTERMEDIAIRES) + 1
entetes_intermediaires(ws, 4, COL_INTER_P, "Proposition : ")
entetes_intermediaires(ws, 4, COL_INTER_S, "Site : ")
ws.column_dimensions[get_column_letter(COL_INTER_P - 1)].width = 3

for i in range(NB_LIGNES_SC):
    r = 5 + i
    exemple = EXEMPLES[i] if i < len(EXEMPLES) else None
    cellule(ws, r, 1, exemple[0] if exemple else None, align=GAUCHE, saisie=True)
    cellule(ws, r, 2, exemple[1] if exemple else None, fmt=FMT_ENTIER, align=DROITE, saisie=True)
    sel = {}
    for j, (cle, _, _) in enumerate(MODULES):
        valeur = None if not exemple else ("Oui" if cle in exemple[2] else "Non")
        c = cellule(ws, r, 3 + j, valeur, align=CENTRE, saisie=True)
        valid_oui_non_sc.add(c)
        sel[cle] = f"${get_column_letter(3 + j)}{r}"
    c = cellule(ws, r, COL_ENG, exemple[3] if exemple else None, fmt="0", align=CENTRE, saisie=True)
    valid_annees_sc.add(c)
    annees_ref = f"${get_column_letter(COL_ENG)}{r}"
    pop_ref = f"$B{r}"
    sel_range = f"$C{r}:${get_column_letter(2 + len(MODULES))}{r}"
    n_expr = f'COUNTIF({sel_range},"Oui")'
    aP = bloc_horizontal(ws, r, COL_INTER_P, "", sel, n_expr, annees_ref, pop_ref)
    aS = bloc_horizontal(ws, r, COL_INTER_S, "Site_", sel, n_expr, annees_ref, pop_ref)
    vide = f'{pop_ref}=""'
    sorties = [
        (f'=IF({vide},"",{aP["mensuel"]})', FMT_EUROS, FILL_RESULTAT, FONT_GRAS),
        (f'=IF({vide},"",{aP["annuel"]})', FMT_EUROS, FILL_RESULTAT, FONT_NORMAL),
        (f'=IF({vide},"",{aP["setup"]})', FMT_EUROS, FILL_RESULTAT, FONT_NORMAL),
        (f'=IF({vide},"",{aP["total"]})', FMT_EUROS, FILL_RESULTAT, FONT_GRAS),
        (f'=IF({vide},"",{aP["seuil"]})', "@", FILL_RESULTAT, FONT_NOTE),
        (f'=IF({vide},"",{aS["mensuel"]})', FMT_EUROS, FILL_SITE, FONT_NORMAL),
        (f'=IF({vide},"",{aS["total"]})', FMT_EUROS, FILL_SITE, FONT_NORMAL),
        (f'=IF({vide},"",{aP["mensuel"]}-{aS["mensuel"]})', FMT_EUROS_SIGNE, None, FONT_NORMAL),
        (f'=IF(OR({vide},{aS["mensuel"]}=0),"",{aP["mensuel"]}/{aS["mensuel"]}-1)', FMT_POURCENT_SIGNE, None, FONT_NORMAL),
    ]
    for k, (formule, fmt, fill, font) in enumerate(sorties):
        cellule(ws, r, COL_SORTIE + k, formule, fmt=fmt, fill=fill, font=font, align=CENTRE if fmt == "@" else DROITE)

plage_ecart = f"{get_column_letter(COL_SORTIE + 7)}5:{get_column_letter(COL_SORTIE + 8)}{4 + NB_LIGNES_SC}"
ws.conditional_formatting.add(plage_ecart, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="B91C1C")))
ws.conditional_formatting.add(plage_ecart, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="15803D")))
ws.sheet_properties.outlinePr.summaryRight = False
proteger(ws)

# ---------------------------------------------------------------------------
# Enregistrement
# ---------------------------------------------------------------------------

wb.calculation.fullCalcOnLoad = True
sortie = Path(__file__).with_name("grille-de-prix-commerciaux.xlsx")
wb.save(sortie)
print(f"Classeur écrit : {sortie}")
